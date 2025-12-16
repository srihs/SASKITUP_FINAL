"""
Word Export Service for Bespoke Quotations

Generates Word documents (ORDER DETAILS format) for quotations containing bespoke products
with player customization details.
"""
from io import BytesIO
from typing import List, Dict
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import logging

logger = logging.getLogger(__name__)


class BespokeQuotationWordGenerator:
    """
    Generates Word document export for quotations with bespoke products and player details.

    Format: ORDER DETAILS document with order header table and product sections.
    """

    # Font configuration
    FONT_NAME = 'Calibri'
    FONT_SIZE_NORMAL = 11
    FONT_SIZE_HEADING = 12

    def __init__(self, quotation):
        """
        Initialize generator with a quotation instance.

        Args:
            quotation: Quotation model instance
        """
        self.quotation = quotation
        self.document = Document()

    def generate(self) -> bytes:
        """
        Generate Word document and return as bytes.

        Returns:
            bytes: Word file content

        Raises:
            ValueError: If quotation has no bespoke items with player customizations
        """
        # Get bespoke items
        bespoke_items = self._get_bespoke_items()

        if not bespoke_items:
            raise ValueError("No player customizations found for bespoke items in this quotation")

        # Set document margins
        self._set_margins()

        # Add title
        self._add_title()

        # Add order details table
        self._add_order_details_table()

        # Add spacing
        self.document.add_paragraph()

        # Group items by base product SKU
        product_groups = self._group_items_by_base_product(bespoke_items)

        # Write all products with page breaks between them
        product_items = list(product_groups.items())
        for index, (base_sku, items) in enumerate(product_items):
            self._write_product_section(base_sku, items)
            # Add page break after each product except the last one
            if index < len(product_items) - 1:
                self.document.add_page_break()

        # Save to BytesIO
        output = BytesIO()
        self.document.save(output)
        output.seek(0)

        return output.getvalue()

    def _set_margins(self):
        """Set document margins (1 inch all sides)"""
        sections = self.document.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

    def _set_table_borderless(self, table):
        """
        Remove all borders from a table.

        Args:
            table: python-docx Table object
        """
        tbl = table._element
        tblPr = tbl.tblPr
        if tblPr is None:
            tblPr = OxmlElement('w:tblPr')
            tbl.insert(0, tblPr)

        # Create table borders element
        tblBorders = OxmlElement('w:tblBorders')
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'none')
            border.set(qn('w:sz'), '0')
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), 'auto')
            tblBorders.append(border)

        tblPr.append(tblBorders)

    def _add_title(self):
        """Add document title"""
        title = self.document.add_paragraph()
        run = title.add_run('ORDER DETAILS')
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.name = self.FONT_NAME
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _add_order_details_table(self):
        """Add order details header table"""
        table = self.document.add_table(rows=5, cols=2)
        table.style = 'Table Grid'

        # Customer name
        table.rows[0].cells[0].text = 'Customer name'
        institution_name = ''
        if self.quotation.institution:
            institution_name = self.quotation.institution.name
        table.rows[0].cells[1].text = institution_name

        # Team / Product Name
        table.rows[1].cells[0].text = 'Team / Product Name'
        table.rows[1].cells[1].text = ''

        # S.O.P No
        table.rows[2].cells[0].text = 'S.O.P No'
        table.rows[2].cells[1].text = self.quotation.quotation_number

        # Ordered Date
        table.rows[3].cells[0].text = 'Ordered Date'
        if self.quotation.created_at:
            table.rows[3].cells[1].text = self.quotation.created_at.strftime('%d %b, %Y')

        # Order required on
        table.rows[4].cells[0].text = 'Order required on'
        if hasattr(self.quotation, 'delivery_date') and self.quotation.delivery_date:
            table.rows[4].cells[1].text = self.quotation.delivery_date.strftime('%d %b, %Y')

        # Style table cells
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = self.FONT_NAME
                        run.font.size = Pt(self.FONT_SIZE_NORMAL)

    def _write_product_section(self, base_sku: str, items: List):
        """
        Write a complete product section (title + all variations + players).

        Args:
            base_sku: Base product SKU
            items: List of QuotationItem instances for this product
        """
        # Get product name from first item
        product_name = items[0].product_name if items else ''

        # Add product title (Heading 5 style)
        product_para = self.document.add_paragraph()
        run = product_para.add_run(f'Product\t\t- {product_name}')
        run.font.size = Pt(self.FONT_SIZE_HEADING)
        run.font.bold = True
        run.font.name = self.FONT_NAME
        product_para.style = 'Heading 5'

        # Add style code
        style_para = self.document.add_paragraph()
        run = style_para.add_run(f'Style Code\t\t- {base_sku}')
        run.font.size = Pt(self.FONT_SIZE_NORMAL)
        run.font.name = self.FONT_NAME

        # Add spacing
        self.document.add_paragraph()

        # Get all variation groups for this product
        variation_groups = self._get_variation_groups_for_product_group(items, base_sku)

        # Write each variation block
        for variation_data in variation_groups:
            self._write_variation_block(variation_data)
            # Add spacing between variations
            self.document.add_paragraph()

    def _write_variation_block(self, variation_data: Dict):
        """
        Write a single variation block (variation info + player table).

        Args:
            variation_data: Dict with variation_code, size, players, total_qty
        """
        # Create variation info table (3 rows x 3 cols) - borderless
        var_table = self.document.add_table(rows=3, cols=3)

        # Remove all borders from variation table
        self._set_table_borderless(var_table)

        # Row 1: Variation style code
        var_table.rows[0].cells[0].text = 'Variation style code'
        var_table.rows[0].cells[1].text = variation_data['variation_code']

        # Row 2: Size
        var_table.rows[1].cells[0].text = 'Size'
        var_table.rows[1].cells[1].text = variation_data['size']

        # Row 3: Total qty
        var_table.rows[2].cells[0].text = 'Total qty'
        var_table.rows[2].cells[1].text = str(variation_data['total_qty'])

        # Style variation table
        for row in var_table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = self.FONT_NAME
                        run.font.size = Pt(self.FONT_SIZE_NORMAL)

        # Add spacing
        self.document.add_paragraph()

        # Create player table
        players = variation_data['players']
        player_table = self.document.add_table(rows=len(players) + 1, cols=3)
        player_table.style = 'Table Grid'

        # Header row
        player_table.rows[0].cells[0].text = 'Name'
        player_table.rows[0].cells[1].text = 'Number'
        player_table.rows[0].cells[2].text = 'Initial'

        # Player rows
        for i, player in enumerate(players, start=1):
            player_table.rows[i].cells[0].text = player.get('player_name', '')
            player_table.rows[i].cells[1].text = str(player.get('player_number', ''))
            player_table.rows[i].cells[2].text = player.get('player_initial', '')

        # Style player table
        for row in player_table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = self.FONT_NAME
                        run.font.size = Pt(self.FONT_SIZE_NORMAL)
                    # Center align header row
                    if row == player_table.rows[0]:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _get_bespoke_items(self) -> List:
        """
        Get all bespoke items from quotation (excluding addons).

        Returns:
            List of QuotationItem instances that are bespoke products
        """
        items = []

        for item in self.quotation.items.filter(is_addon=False).select_related('product_content_type'):
            # Check if this is a bespoke product
            if item.product_content_type.model.lower() == 'bespokeproduct':
                # Check if it has player customizations
                player_customizations = item.variations.get('player_customizations', [])
                if player_customizations:
                    items.append(item)

        return items

    def _group_items_by_base_product(self, bespoke_items):
        """
        Group QuotationItems by base product SKU.

        Args:
            bespoke_items: List of QuotationItem instances

        Returns:
            dict: {base_sku: [item1, item2, ...]}
        """
        groups = {}
        for item in bespoke_items:
            # Extract base SKU (remove size suffix)
            base_sku = item.product_sku or ''
            if base_sku and '-' in base_sku:
                base_sku = base_sku.rsplit('-', 1)[0].strip()

            if base_sku not in groups:
                groups[base_sku] = []
            groups[base_sku].append(item)

        return groups

    def _get_variation_groups_for_product_group(self, items: List, base_sku: str) -> List[Dict]:
        """
        Group players by variation (size) across ALL items for a base product.

        Args:
            items: List of QuotationItem instances for the same base product
            base_sku: Base product SKU (without size suffix)

        Returns:
            List of dicts with variation data and players grouped by size
        """
        # Collect all players from all items
        all_players = []
        for item in items:
            player_customizations = item.variations.get('player_customizations', [])
            all_players.extend(player_customizations)

        if not all_players:
            return []

        # Group players by size
        size_groups = {}
        for player in all_players:
            size = player.get('size', '') or player.get('color', '') or 'N/A'

            if size not in size_groups:
                size_groups[size] = []

            size_groups[size].append(player)

        # Convert to list of variation groups
        variation_groups = []
        for size, players in size_groups.items():
            # Generate variation code (base_sku-size)
            variation_code = f"{base_sku}-{size}" if base_sku else size

            variation_groups.append({
                'variation_code': variation_code,
                'size': size,
                'players': players,
                'total_qty': len(players)
            })

        return variation_groups


def generate_bespoke_quotation_word(quotation) -> bytes:
    """
    Convenience function to generate Word document for a quotation.

    Args:
        quotation: Quotation model instance

    Returns:
        bytes: Word file content
    """
    generator = BespokeQuotationWordGenerator(quotation)
    return generator.generate()
